# -*- coding: utf-8 -*-
"""Разбор Excel для массового создания ивентов Carrot + резолв Media."""

from __future__ import annotations

import io
import os
import re
from pathlib import Path
from typing import Any, Callable, Optional

from openpyxl import load_workbook

HEADER_NAME_ALIASES = {
    "name",
    "имя",
    "название",
    "event",
    "event name",
    "имя ивента",
    "имя события",
}

_GUID_RE = re.compile(
    r"^[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}$"
)


def _cell_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def is_media_type(type_name: Any) -> bool:
    return str(type_name or "").strip().lower() == "media"


def read_excel_rows(file_bytes: bytes) -> list[list[str]]:
    """
    Прочитать первый лист .xlsx/.xlsm как матрицу строк.

    Пустые строки в конце отбрасываются. Ячейки → строки.
    """
    if not file_bytes:
        raise ValueError("Пустой Excel-файл")

    try:
        wb = load_workbook(
            filename=io.BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )
    except Exception as err:  # noqa: BLE001
        raise ValueError(
            "Не удалось открыть Excel. Нужен .xlsx/.xlsm (не старый .xls)."
        ) from err

    try:
        ws = wb.active
        if ws is None:
            raise ValueError("В книге нет активного листа")
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            cells = [_cell_str(c) for c in (row or ())]
            while cells and cells[-1] == "":
                cells.pop()
            if not any(cells):
                continue
            rows.append(cells)
    finally:
        wb.close()

    if not rows:
        raise ValueError("В Excel нет данных")
    return rows


def _looks_like_header(row: list[str]) -> bool:
    if not row:
        return False
    first = (row[0] or "").strip().lower()
    return first in HEADER_NAME_ALIASES


def rows_to_event_drafts(
    rows: list[list[str]],
    template: dict[str, Any],
    *,
    has_header: Optional[bool] = None,
) -> dict[str, Any]:
    """
    Преобразовать строки Excel в черновики ивентов.

    Колонки:
      1 — имя ивента
      2..N-1 — значения переменных шаблона (по порядку template.variables)
      N — External id

    Для Media-переменных в середине ожидается имя файла (не id;name).
    """
    if not rows:
        raise ValueError("Нет строк для импорта")

    skip_header = has_header
    if skip_header is None:
        skip_header = _looks_like_header(rows[0])

    data_rows = rows[1:] if skip_header else rows
    var_meta = list(template.get("variables") or [])
    template_id = template.get("id") or ""
    if not template_id:
        raise ValueError("У шаблона нет id")

    drafts: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []

    for idx, row in enumerate(data_rows):
        excel_row = idx + (2 if skip_header else 1)
        if len(row) < 2:
            skipped.append(
                {
                    "row": excel_row,
                    "reason": "нужны минимум 2 столбца (имя и External id)",
                }
            )
            continue

        name = row[0]
        external_id = row[-1]
        mid = row[1:-1]

        if not name:
            skipped.append({"row": excel_row, "reason": "пустое имя"})
            continue

        variables: list[dict[str, Any]] = []
        for vi, meta in enumerate(var_meta):
            nv = dict(meta)
            if vi < len(mid) and mid[vi] != "":
                nv["value"] = mid[vi]
            else:
                nv["value"] = (
                    meta.get("value")
                    if meta.get("value") not in (None, "")
                    else (meta.get("defaultValue") or "")
                )
            variables.append(nv)

        drafts.append(
            {
                "name": name,
                "externalId": external_id,
                "templateId": template_id,
                "contentId": template.get("contentId") or "",
                "templateTypeInt": template.get("templateTypeInt"),
                "state": template.get("defaultInState") or "IN",
                "comment": "",
                "variables": variables,
                "_excelRow": excel_row,
            }
        )

    return {
        "drafts": drafts,
        "skipped": skipped,
        "headerSkipped": bool(skip_header),
        "variableCount": len(var_meta),
        "variableNames": [v.get("name") or "" for v in var_meta],
        "mediaVarIndexes": [
            i for i, v in enumerate(var_meta) if is_media_type(v.get("type"))
        ],
    }


def _norm_name(name: str) -> str:
    return os.path.basename(str(name or "").strip()).lower()


def _parse_media_value(raw: str) -> tuple[str, str]:
    s = str(raw or "").strip()
    if not s:
        return "", ""
    if ";" in s:
        aid, aname = s.split(";", 1)
        return aid.strip(), aname.strip()
    return s, ""


def index_assets_by_name(assets: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    """Индекс ассетов по имени файла (lower) и stem без расширения."""
    index: dict[str, dict[str, Any]] = {}
    for asset in assets or []:
        name = str(asset.get("name") or "").strip()
        if not name:
            continue
        key = _norm_name(name)
        index.setdefault(key, asset)
        stem = Path(key).stem
        if stem and stem not in index:
            index[stem] = asset
    return index


def find_local_media_file(folder: str, filename: str) -> Optional[Path]:
    """
    Найти файл в папке по имени (с расширением или без).

    Только внутри folder (path traversal блокируется).
    """
    root = Path(folder).expanduser()
    if not root.is_dir():
        raise ValueError(f"Папка медиа не найдена: {folder}")

    root_resolved = root.resolve()
    wanted = os.path.basename(str(filename or "").strip())
    if not wanted or wanted in (".", ".."):
        return None

    candidate = (root / wanted).resolve()
    try:
        candidate.relative_to(root_resolved)
    except ValueError as err:
        raise ValueError("Имя файла выходит за пределы папки медиа") from err

    if candidate.is_file():
        return candidate

    # без расширения — первое совпадение по stem
    stem = Path(wanted).stem.lower()
    if not stem:
        return None
    for path in sorted(root_resolved.iterdir()):
        if path.is_file() and path.stem.lower() == stem:
            return path
    return None


def resolve_media_in_drafts(
    drafts: list[dict[str, Any]],
    *,
    assets: list[dict[str, Any]],
    media_folder: Optional[str] = None,
    upload_file: Optional[Callable[[bytes, str], dict[str, Any]]] = None,
) -> dict[str, Any]:
    """
    Для Media-переменных: значение из Excel = имя файла.

    1) если уже «guid;name» — оставить
    2) если ассет с таким именем есть в Carrot — подставить id;name
    3) иначе искать в media_folder и upload_file(...)
    """
    index = index_assets_by_name(assets)
    resolved = 0
    uploaded = 0
    reused = 0
    missing: list[dict[str, Any]] = []
    folder = (media_folder or "").strip() or None

    for draft in drafts:
        excel_row = draft.get("_excelRow")
        for var in draft.get("variables") or []:
            if not is_media_type(var.get("type")):
                continue
            raw = "" if var.get("value") is None else str(var.get("value")).strip()
            if not raw:
                continue

            aid, aname = _parse_media_value(raw)
            if aid and _GUID_RE.match(aid):
                # уже готовое значение Carrot
                if aname:
                    var["value"] = f"{aid};{aname}"
                else:
                    var["value"] = aid
                resolved += 1
                continue

            filename = aname or aid or raw
            key = _norm_name(filename)
            hit = index.get(key) or index.get(Path(key).stem)
            if hit and hit.get("id"):
                var["value"] = f"{hit['id']};{hit.get('name') or filename}"
                resolved += 1
                reused += 1
                continue

            if not folder:
                missing.append(
                    {
                        "row": excel_row,
                        "variable": var.get("name"),
                        "file": filename,
                        "reason": "медиа нет в Carrot; укажи папку в Настройках",
                    }
                )
                continue

            if upload_file is None:
                missing.append(
                    {
                        "row": excel_row,
                        "variable": var.get("name"),
                        "file": filename,
                        "reason": "нет обработчика загрузки",
                    }
                )
                continue

            try:
                local = find_local_media_file(folder, filename)
            except ValueError as err:
                missing.append(
                    {
                        "row": excel_row,
                        "variable": var.get("name"),
                        "file": filename,
                        "reason": str(err),
                    }
                )
                continue

            if local is None:
                missing.append(
                    {
                        "row": excel_row,
                        "variable": var.get("name"),
                        "file": filename,
                        "reason": f"файл не найден в папке {folder}",
                    }
                )
                continue

            asset = upload_file(local.read_bytes(), local.name)
            index[_norm_name(asset.get("name") or local.name)] = asset
            stem = Path(_norm_name(asset.get("name") or local.name)).stem
            if stem:
                index.setdefault(stem, asset)
            var["value"] = f"{asset['id']};{asset.get('name') or local.name}"
            resolved += 1
            uploaded += 1

    return {
        "resolved": resolved,
        "uploaded": uploaded,
        "reused": reused,
        "missing": missing,
        # совместимость со старым UI
        "warnings": missing,
        "ok": not missing,
    }
